"""
Імпорт запчастин з XLSX у warehouse.db + завантаження фото через Telegram API.

Використання (бот може бути запущений або ні):
    python import_xlsx.py <xlsx> <TOKEN> <YOUR_USER_ID>

Приклад:
    python import_xlsx.py warehouse.xlsx "123456:AAF..." 655446281
"""

import sys, os, io, sqlite3, time, requests, zipfile
import xml.etree.ElementTree as ET
import openpyxl
from collections import defaultdict

NS = {
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'a':   'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r':   'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}
R_EMBED = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'


def get_image_map(xlsx_path: str, sheet_index: int) -> dict:
    """Extract row(0-based) -> image bytes from xlsx drawing XML.
    Correctly maps sheet to its drawing via worksheet rels file."""
    row_to_img = {}

    with zipfile.ZipFile(xlsx_path) as z:
        # Find the correct drawing for this sheet via _rels
        ws_rels_path = f'xl/worksheets/_rels/sheet{sheet_index + 1}.xml.rels'
        draw_path = None
        if ws_rels_path in z.namelist():
            rels_root = ET.fromstring(z.read(ws_rels_path))
            for rel in rels_root:
                tgt = rel.attrib.get('Target', '')
                if 'drawing' in tgt.lower():
                    draw_path = tgt.replace('../', 'xl/')
                    break

        if not draw_path:
            return {}

        rels_path = draw_path.replace('xl/drawings/', 'xl/drawings/_rels/').replace('.xml', '.xml.rels')
        if draw_path not in z.namelist() or rels_path not in z.namelist():
            return {}

        # rId -> media file path
        rels_root = ET.fromstring(z.read(rels_path).decode('utf-8'))
        rid_to_file = {}
        for rel in rels_root:
            rid = rel.attrib['Id']
            tgt = rel.attrib['Target'].replace('../', 'xl/')
            rid_to_file[rid] = tgt

        # row -> rId
        draw_root = ET.fromstring(z.read(draw_path).decode('utf-8'))
        for anchor in draw_root.findall('xdr:oneCellAnchor', NS):
            frm = anchor.find('xdr:from', NS)
            row = int(frm.find('xdr:row', NS).text)
            pic = anchor.find('xdr:pic', NS)
            if pic is None:
                continue
            blip = pic.find('xdr:blipFill/a:blip', NS)
            if blip is None:
                continue
            rid = blip.attrib.get(R_EMBED)
            if not rid:
                continue
            media = rid_to_file.get(rid)
            if media and media in z.namelist() and row not in row_to_img:
                row_to_img[row] = z.read(media)

    return row_to_img


def parse_xlsx(path: str) -> list:
    wb = openpyxl.load_workbook(path)
    merged = defaultdict(lambda: {
        'part_number': '', 'car_brand': '', 'car_model': '',
        'description': '', 'price': None, 'quantity': 0,
        'condition': set(), 'market': None,
        'side': None, 'lamp_type': None, 'img_bytes': None,
    })

    for sheet_idx, sname in enumerate(wb.sheetnames):
        ws = wb[sname]

        # Get image map via XML (works even when openpyxl._images is empty)
        img_map = get_image_map(path, sheet_idx)
        print(f"  Аркуш '{sname}': рядків={ws.max_row}, фото={len(img_map)}")

        # Auto-detect header row
        header_row = 2
        for r in range(1, 6):
            if str(ws.cell(r, 2).value or '').strip() == 'Номер':
                header_row = r
                break

        # Auto-detect column indices by header name
        col_map = {}
        for c in range(1, ws.max_column + 1):
            hdr = str(ws.cell(header_row, c).value or '').strip().lower()
            col_map[hdr] = c

        # Default columns (fallback if header not found)
        COL_NUM    = col_map.get('номер',       2)
        COL_BRAND  = col_map.get('марка авто',  3)
        COL_MODEL  = col_map.get('модель авто', 4)
        COL_DESC   = col_map.get('опис',        5)
        COL_PRICE  = col_map.get('ціна',        7)
        COL_QTY    = col_map.get('кількість',   9)
        COL_NEW    = col_map.get('нова',        11)
        COL_USED   = col_map.get('б/в',         12)
        COL_REPAIR = col_map.get('до роб.',     14)
        COL_BROKEN = col_map.get('битий',       15)
        COL_EU     = col_map.get('europa',      16)
        COL_USA    = col_map.get('usa',         17)
        COL_SIDE   = col_map.get('сторона',     21)  # U
        COL_LAMP   = col_map.get('тип фари',    22)  # V

        print(f"    Колонки: сторона={COL_SIDE}, тип фари={COL_LAMP}")

        for excel_row in range(header_row + 1, ws.max_row + 1):
            def cell(c): return ws.cell(excel_row, c).value
            num = str(cell(COL_NUM) or '').strip().upper()
            if not num:
                continue

            brand = str(cell(COL_BRAND) or '').strip()
            model = str(cell(COL_MODEL) or '').strip()
            desc  = str(cell(COL_DESC)  or '').strip()

            def is_true(c):
                return str(cell(c) or '').strip().upper() in (
                    'TRUE', '1', 'ІСТИНА', 'ПРАВДА', 'ИСТИНА'
                )

            key = (num, brand.lower(), model.lower())
            m   = merged[key]
            m['part_number'] = num
            m['car_brand']   = brand
            m['car_model']   = model
            m['description'] = m['description'] or desc

            try:
                p = float(cell(COL_PRICE))
                if m['price'] is None or p > m['price']:
                    m['price'] = p
            except (TypeError, ValueError):
                pass

            try:
                m['quantity'] += int(cell(COL_QTY))
            except (TypeError, ValueError):
                pass

            if is_true(COL_NEW):    m['condition'].add('Нова')
            if is_true(COL_USED):   m['condition'].add('Б/В')
            if is_true(COL_REPAIR): m['condition'].add('До роботи')
            if is_true(COL_BROKEN): m['condition'].add('Битий')
            if   is_true(COL_EU) and is_true(COL_USA): m['market'] = 'EU, USA'
            elif is_true(COL_EU):  m['market'] = m['market'] or 'EU'
            elif is_true(COL_USA): m['market'] = m['market'] or 'USA'

            # Side (Сторона): text value
            side_val = str(cell(COL_SIDE) or '').strip()
            if side_val and not m['side']:
                m['side'] = side_val

            # Lamp type (Тип фари): text value
            lamp_val = str(cell(COL_LAMP) or '').strip()
            if lamp_val and not m['lamp_type']:
                m['lamp_type'] = lamp_val

            # anchor_row (0-based) == excel_row - 1
            anchor_r = excel_row - 1
            if m['img_bytes'] is None and anchor_r in img_map:
                m['img_bytes'] = img_map[anchor_r]

    parts = list(merged.values())
    for p in parts:
        p['condition'] = ', '.join(sorted(p['condition'])) or None

    with_photo = sum(1 for p in parts if p['img_bytes'])
    print(f"\n  Унікальних: {len(parts)}, з фото: {with_photo}")
    return parts


def upload_photo(token: str, chat_id: str, img_bytes: bytes) -> str | None:
    url = f"https://api.telegram.org/bot{token}/sendPhoto"
    try:
        resp = requests.post(
            url,
            data={'chat_id': chat_id},
            files={'photo': ('photo.jpg', io.BytesIO(img_bytes), 'image/jpeg')},
            timeout=30
        )
        data = resp.json()
        if not data.get('ok'):
            return None
        msg_id  = data['result']['message_id']
        file_id = data['result']['photo'][-1]['file_id']
        requests.post(
            f"https://api.telegram.org/bot{token}/deleteMessage",
            data={'chat_id': chat_id, 'message_id': msg_id},
            timeout=10
        )
        return file_id
    except Exception as e:
        return None


def import_to_db(parts: list, token: str, chat_id: str, db_path: str = 'warehouse.db'):
    if not os.path.exists(db_path):
        print(f"❌ БД не знайдена: {db_path}. Спочатку запустіть бота хоча б раз.")
        sys.exit(1)

    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    cols = [r[1] for r in conn.execute("PRAGMA table_info(parts)").fetchall()]
    for col in ['market', 'side', 'lamp_type']:
        if col not in cols:
            conn.execute(f"ALTER TABLE parts ADD COLUMN {col} TEXT")

    inserted = updated = skipped = photo_ok = photo_fail = 0

    print(f"\n⏳ Імпортую {len(parts)} запчастин...")

    for i, p in enumerate(parts, 1):
        if not p['part_number'] or not p['car_brand']:
            skipped += 1
            continue

        photo_id = None
        if p['img_bytes']:
            photo_id = upload_photo(token, chat_id, p['img_bytes'])
            if photo_id:
                photo_ok += 1
            else:
                photo_fail += 1
            time.sleep(0.05)

        exists = conn.execute(
            "SELECT id FROM parts WHERE UPPER(part_number)=? AND UPPER(car_brand)=?",
            (p['part_number'].upper(), p['car_brand'].upper())
        ).fetchone()

        if exists:
            conn.execute(
                """UPDATE parts SET
                   quantity   = quantity + ?,
                   photo_id   = COALESCE(photo_id, ?),
                   updated_at = CURRENT_TIMESTAMP
                   WHERE id = ?""",
                (p['quantity'], photo_id, exists['id'])
            )
            updated += 1
        else:
            conn.execute(
                """INSERT INTO parts
                   (part_number, car_brand, car_model, description,
                    price, quantity, condition, market, side, lamp_type, photo_id)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                (p['part_number'], p['car_brand'], p['car_model'] or None,
                 p['description'] or None, p['price'], p['quantity'],
                 p['condition'], p['market'], p['side'], p['lamp_type'],
                 photo_id)
            )
            inserted += 1

        if i % 25 == 0:
            conn.commit()
            print(f"   {i}/{len(parts)}... (фото: ✅{photo_ok} ❌{photo_fail})")

    conn.commit()
    conn.close()

    print(f"\n✅ Готово!")
    print(f"   ➕ Додано:    {inserted}")
    print(f"   🔄 Оновлено:  {updated}")
    print(f"   ⏭  Пропущено: {skipped}")
    print(f"   🖼  Фото:  ✅{photo_ok}  ❌{photo_fail}")


def main():
    if len(sys.argv) < 4:
        print("Використання: python import_xlsx.py <xlsx> <TOKEN> <USER_ID>")
        print('Приклад:      python import_xlsx.py warehouse.xlsx "123:AAF..." 655446281')
        print("\nВаш ID: напишіть @userinfobot у Telegram")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    token     = sys.argv[2]
    chat_id   = sys.argv[3]

    if not os.path.exists(xlsx_path):
        print(f"❌ Файл не знайдено: {xlsx_path}")
        sys.exit(1)

    resp = requests.get(f"https://api.telegram.org/bot{token}/getMe", timeout=10)
    if not resp.json().get('ok'):
        print("❌ Невірний токен або немає інтернету")
        sys.exit(1)
    print(f"🤖 Бот: @{resp.json()['result']['username']}")
    print(f"📨 Чат: {chat_id}")

    print(f"\n📂 Читаю {xlsx_path}...")
    parts = parse_xlsx(xlsx_path)
    import_to_db(parts, token, chat_id)


if __name__ == '__main__':
    main()
